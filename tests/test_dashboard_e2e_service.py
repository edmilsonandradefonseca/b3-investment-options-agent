from datetime import date

import openpyxl

from b3_agent.dashboard_e2e import DashboardE2EService


def _portfolio(path):
    wb = openpyxl.Workbook()
    capa = wb.active
    capa.title = "Capa"
    capa["C1"] = "Período de 27/09/26 a 27/09/26"
    ws = wb.create_sheet("Renda Variavel")
    ws.append(["", "Posição > Ações"])
    ws.append(["", "Código", "Ação", "Qtde.", "Preço Fechamento R$", "Preço Médio R$", "Saldo Bruto R$"])
    ws.append(["", "PETR4*", "PETROBRAS", 1000, 30.0, 28.0, 30000.0])
    ws.append(["", "Total em Ações R$", "", "", "", "", 30000.0])
    ws.append(["", "Posição > Opções"])
    ws.append(["", "Código", "Ativo Ref.", "Qtde.", "Preço Exercício R$", "Data Exercício R$", "Tipo", "Posição", "Prêmio Mercado R$", "Valor de Mercado R$"])
    ws.append(["", "PETRJ360*", "PETR4", -10, 36.0, date(2026, 10, 16), "Call", "vendida", 0.5, -500.0])
    ws.append(["", "Total em Opções R$", "", "", "", "", "", "", "", -500.0])
    wb.save(path)


def _options(path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["Ativo", "Corretora", "Qtd.", "Custo Médio", "Custo Total"])
    ws.append(["PETRJ360", "BTG", -10, 0.4, -400.0])
    wb.save(path)


def test_realistic_excel_to_dashboard_snapshot_is_deterministic(tmp_path):
    portfolio = tmp_path / "portfolio.xlsx"
    options = tmp_path / "options.xlsx"
    _portfolio(portfolio)
    _options(options)

    snapshot = DashboardE2EService().load(portfolio, options_path=options)

    assert snapshot.portfolio.as_of == date(2026, 9, 27)
    assert len(snapshot.portfolio.positions) == 2
    assert len(snapshot.option_transactions) == 1
    assert snapshot.portfolio_intelligence.capital_risk is not None
    assert snapshot.opportunities.as_of == date(2026, 9, 27)
    assert snapshot.opportunities.ranked_opportunities == ()
    assert snapshot.opportunities.rejected_opportunities == ()
    assert snapshot.opportunities.source_refs == ("BTG:Renda Variavel",)


def test_options_snapshot_is_optional(tmp_path):
    portfolio = tmp_path / "portfolio.xlsx"
    _portfolio(portfolio)

    snapshot = DashboardE2EService().load(portfolio)

    assert snapshot.option_transactions == ()
    assert len(snapshot.portfolio.positions) == 2
