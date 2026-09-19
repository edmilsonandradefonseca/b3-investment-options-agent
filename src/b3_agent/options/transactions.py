from __future__ import annotations

from pathlib import Path

from openpyxl import load_workbook

from b3_agent.schemas.option_transaction import OptionTransaction


class OptionTransactionIngestionError(ValueError):
    """Raised when the options transaction workbook cannot be parsed."""


def _text(value: object) -> str:
    return "" if value is None else str(value).strip().replace("*", "")


def _number(value: object) -> float | None:
    if value is None or (isinstance(value, str) and value.strip() in {"", "-"}):
        return None
    try:
        return float(value)
    except (TypeError, ValueError) as exc:
        raise OptionTransactionIngestionError(f"expected numeric value, got {value!r}") from exc


class OptionsTransactionLoader:
    """Load the user's options transaction history from the operations workbook.

    The loader preserves the workbook semantics: `Custo Médio` and `Custo Total`
    remain transaction-cost fields and are not relabeled as market premium.
    """

    def load(self, path: str | Path) -> tuple[OptionTransaction, ...]:
        workbook = load_workbook(Path(path), data_only=True, read_only=True)
        try:
            sheet = workbook.active
            rows = list(sheet.iter_rows(values_only=True))
            if not rows:
                raise OptionTransactionIngestionError("transaction workbook is empty")

            header = {str(value).strip(): index for index, value in enumerate(rows[0]) if value is not None}
            required = {"Ativo", "Corretora", "Qtd.", "Custo Médio", "Custo Total"}
            missing = required - header.keys()
            if missing:
                raise OptionTransactionIngestionError(f"missing columns: {sorted(missing)}")

            transactions: list[OptionTransaction] = []
            for index, row in enumerate(rows[1:], start=2):
                ticker = _text(row[header["Ativo"]])
                if not ticker:
                    continue
                quantity = _number(row[header["Qtd."]])
                if quantity in (None, 0):
                    continue
                transactions.append(
                    OptionTransaction(
                        transaction_id=f"options-xlsx:{index}:{ticker}",
                        option_ticker=ticker,
                        broker=_text(row[header["Corretora"]]),
                        quantity=quantity,
                        average_cost=_number(row[header["Custo Médio"]]),
                        total_cost=_number(row[header["Custo Total"]]),
                        source_ref="Options Transactions XLSX",
                        source_type="OPTIONS_XLSX",
                        source_id=Path(path).name,
                    )
                )
            return tuple(transactions)
        finally:
            workbook.close()
