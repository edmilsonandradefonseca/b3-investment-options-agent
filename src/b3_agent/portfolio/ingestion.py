from __future__ import annotations

import re
from datetime import date, datetime
from pathlib import Path

from openpyxl import load_workbook

from b3_agent.schemas.position import PortfolioContext, Position


class PortfolioIngestionError(ValueError):
    """Raised when the authoritative BTG portfolio cannot be parsed."""


def _text(value: object) -> str:
    return "" if value is None else str(value).strip().replace("*", "")


def _number(value: object) -> float | None:
    if value is None or (isinstance(value, str) and value.strip() in {"", "-"}):
        return None
    try:
        return float(value)
    except (TypeError, ValueError) as exc:
        raise PortfolioIngestionError(f"expected numeric value, got {value!r}") from exc


def _as_date(value: object) -> date:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    raise PortfolioIngestionError(f"expected date, got {value!r}")


def _find_section(rows: list[tuple[object, ...]], label: str) -> int:
    for index, row in enumerate(rows):
        if any(isinstance(value, str) and label in value for value in row):
            return index
    raise PortfolioIngestionError(f"section not found: {label}")


def _statement_date(workbook) -> date:
    for row in workbook["Capa"].iter_rows(values_only=True):
        for value in row:
            if isinstance(value, str):
                match = re.search(r"Período de .* a (\d{2}/\d{2}/\d{2})", value)
                if match:
                    return datetime.strptime(match.group(1), "%d/%m/%y").date()
    raise PortfolioIngestionError("statement period end not found")


class BtgRendaVariavelLoader:
    """Load current positions from BTG's authoritative `Renda Variavel` sheet."""

    def load(self, path: str | Path) -> PortfolioContext:
        workbook = load_workbook(Path(path), data_only=True, read_only=True)
        if "Renda Variavel" not in workbook.sheetnames:
            raise PortfolioIngestionError("Renda Variavel sheet not found")
        rows = list(workbook["Renda Variavel"].iter_rows(values_only=True))
        as_of = _statement_date(workbook)
        positions: list[Position] = []

        actions_start = _find_section(rows, "Posição > Ações")
        for row in rows[actions_start + 2 :]:
            ticker = _text(row[1])
            if ticker.startswith("Total em Ações"):
                break
            if not ticker or ticker in {"Código", "Posição"} or _number(row[3]) in (None, 0):
                continue
            positions.append(Position(
                position_id=f"btg:renda-variavel:{ticker}", ticker=ticker,
                instrument_type="STOCK", quantity=_number(row[3]),
                average_cost=_number(row[5]), market_price=_number(row[4]),
                market_value=_number(row[6]), source_ref="BTG:Renda Variavel:Acoes",
            ))

        options_start = _find_section(rows, "Posição > Opções")
        for row in rows[options_start + 2 :]:
            ticker = _text(row[1])
            if ticker.startswith("Total em Opções"):
                break
            if not ticker or ticker == "Código" or _number(row[3]) in (None, 0):
                continue
            option_type = _text(row[6]).upper()
            if option_type not in {"PUT", "CALL"}:
                raise PortfolioIngestionError(f"unsupported option type {row[6]!r} for {ticker}")
            positions.append(Position(
                position_id=f"btg:renda-variavel:{ticker}", ticker=ticker,
                instrument_type="OPTION", quantity=_number(row[3]),
                strike=_number(row[4]), expiration_date=_as_date(row[5]),
                option_type=option_type, underlying_ticker=_text(row[2]),
                contract_multiplier=1.0, market_price=_number(row[8]),
                market_value=_number(row[9]), source_ref="BTG:Renda Variavel:Opcoes",
            ))

        return PortfolioContext(
            as_of=as_of, positions=tuple(positions), cash=0.0,
            source_refs=("BTG:Renda Variavel",), quality_status="VALIDATED",
        )
